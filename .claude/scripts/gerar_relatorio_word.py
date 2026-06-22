#!/usr/bin/env python3
"""
Gerador de Relatório de Andamento Processual em Word (.docx)
Escritório Flaviane Bilhar Caler — Direito da Saúde & Previdenciário
"""

import sys
import json
import re
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


INDIGO = RGBColor(0x31, 0x2e, 0x81)   # indigo-900
DARK   = RGBColor(0x1a, 0x1a, 0x1a)
GRAY   = RGBColor(0x55, 0x55, 0x55)
AMBER  = RGBColor(0x92, 0x40, 0x08)


def set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_hrule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "312e81")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def build_relatorio(dados: dict, secoes: dict, output_path: str):
    doc = Document()

    # Margens
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(2.5)

    # ── CABEÇALHO ──────────────────────────────────────────────
    cab = doc.add_paragraph()
    cab.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cab.add_run("ESCRITÓRIO FLAVIANE BILHAR CALER")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = INDIGO

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("Advogada | OAB/SC   ·   Direito da Saúde & Direito Previdenciário")
    rs.font.size = Pt(9)
    rs.font.color.rgb = GRAY

    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs2 = sub2.add_run("Chapecó/SC")
    rs2.font.size = Pt(9)
    rs2.font.color.rgb = GRAY

    add_hrule(doc)

    # Título
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_before = Pt(12)
    rt = titulo.add_run("RELATÓRIO DE ANDAMENTO PROCESSUAL")
    rt.bold = True
    rt.font.size = Pt(13)
    rt.font.color.rgb = DARK

    data_p = doc.add_paragraph()
    data_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    data_p.paragraph_format.space_after = Pt(10)
    rd = data_p.add_run(f"Data: {datetime.now().strftime('%d/%m/%Y')}")
    rd.font.size = Pt(9)
    rd.font.color.rgb = GRAY

    add_hrule(doc)

    # ── 1. IDENTIFICAÇÃO ────────────────────────────────────────
    _secao_titulo(doc, "1. IDENTIFICAÇÃO")

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    campos = [
        ("Cliente",           dados.get("cliente", "")),
        ("Processo nº",       dados.get("processo", "")),
        ("Vara / Juízo",      dados.get("vara", "")),
        ("Parte Contrária",   dados.get("reu", "")),
        ("Tipo de Ação",      dados.get("tipo_acao", "")),
        ("Fase Processual",   dados.get("fase", "")),
    ]
    for label, valor in campos:
        row = table.add_row()
        row.cells[0].width = Cm(5)
        row.cells[1].width = Cm(12)
        set_cell_bg(row.cells[0], "f3f4f6")
        _cell_text(row.cells[0], label, bold=True, size=9)
        _cell_text(row.cells[1], valor, size=9)

    doc.add_paragraph()

    # ── SEÇÕES NARRATIVAS ───────────────────────────────────────
    ordem = [
        ("2. SITUAÇÃO PROCESSUAL ATUAL",    "situacao_atual"),
        ("3. ÚLTIMA MOVIMENTAÇÃO",          "ultima_movimentacao"),
        ("4. PRÓXIMAS ETAPAS",              "proximas_etapas"),
        ("5. PERSPECTIVA JURÍDICA",         "perspectiva_juridica"),
        ("6. ORIENTAÇÕES AO CLIENTE",       "orientacoes"),
        ("7. OBSERVAÇÕES FINAIS",           "observacoes_finais"),
    ]
    for titulo_sec, chave in ordem:
        texto = secoes.get(chave, "")
        if not texto:
            continue
        _secao_titulo(doc, titulo_sec)
        _corpo(doc, texto)

    # ── AVISO (documentos necessários) ─────────────────────────
    doc_necessario = dados.get("doc_necessario", "").strip()
    if doc_necessario:
        add_hrule(doc)
        aviso_p = doc.add_paragraph()
        aviso_p.paragraph_format.space_before = Pt(6)
        ra = aviso_p.add_run("⚠  DOCUMENTOS NECESSÁRIOS: ")
        ra.bold = True
        ra.font.color.rgb = AMBER
        ra.font.size = Pt(10)
        aviso_p.add_run(doc_necessario).font.size = Pt(10)

    # ── RODAPÉ ─────────────────────────────────────────────────
    add_hrule(doc)
    doc.add_paragraph()

    rod = doc.add_paragraph()
    rod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = rod.add_run("Dra. Flaviane Bilhar Caler\nOAB/SC — Direito da Saúde & Previdenciário")
    rr.bold = True
    rr.font.size = Pt(10)
    rr.font.color.rgb = INDIGO

    rod2 = doc.add_paragraph()
    rod2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr2 = rod2.add_run(
        "Este relatório é de uso exclusivo do cliente identificado acima e não "
        "constitui parecer jurídico público.\nPara dúvidas, entre em contato com o escritório."
    )
    rr2.font.size = Pt(8)
    rr2.font.color.rgb = GRAY

    doc.save(output_path)
    print(f"✓ Relatório salvo em: {output_path}")


def _secao_titulo(doc, texto: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(texto)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = INDIGO

    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "d1d5db")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _corpo(doc, texto: str):
    linhas = texto.strip().split("\n")
    for linha in linhas:
        linha = linha.strip()
        if not linha:
            continue
        # Itens de lista
        if linha.startswith("- ") or linha.startswith("• "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            p.add_run(linha[2:]).font.size = Pt(10)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            # Negrito inline **texto**
            partes = re.split(r"\*\*(.+?)\*\*", linha)
            for i, parte in enumerate(partes):
                r = p.add_run(parte)
                r.font.size = Pt(10)
                if i % 2 == 1:
                    r.bold = True


def _cell_text(cell, texto: str, bold=False, size=10):
    cell.paragraphs[0].clear()
    r = cell.paragraphs[0].add_run(texto)
    r.bold = bold
    r.font.size = Pt(size)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python3 gerar_relatorio_word.py <dados.json>")
        sys.exit(1)

    with open(sys.argv[1], encoding="utf-8") as f:
        payload = json.load(f)

    dados  = payload.get("dados", {})
    secoes = payload.get("secoes", {})
    nome_arquivo = dados.get("cliente", "cliente").replace(" ", "_")
    output = payload.get("output", f"Relatorio_{nome_arquivo}.docx")

    build_relatorio(dados, secoes, output)
